import os
import sys
import json
import urllib.request
import urllib.parse
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document

BASE_URL = "http://127.0.0.1:8000"

def post_json(endpoint: str, payload: dict):
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}{endpoint}",
        data=data,
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as response:
        return json.loads(response.read().decode("utf-8"))

def get_binary(url_path: str):
    full_url = f"{BASE_URL}{url_path}" if url_path.startswith("/") else url_path
    req = urllib.request.Request(full_url)
    with urllib.request.urlopen(req) as response:
        return response.headers.get("content-type"), response.read()

def post_multipart_file(endpoint: str, file_path: Path, filename: str):
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    with open(file_path, "rb") as f:
        file_bytes = f.read()

    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="resume"; filename="{filename}"\r\n'
        f"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
    ).encode("utf-8") + file_bytes + f"\r\n--{boundary}--\r\n".encode("utf-8")

    req = urllib.request.Request(
        f"{BASE_URL}{endpoint}",
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}
    )
    with urllib.request.urlopen(req) as response:
        return json.loads(response.read().decode("utf-8"))

def test_live_backend():
    print("==================================================", flush=True)
    print("Testing LIVE Backend Server at http://127.0.0.1:8000...", flush=True)
    print("==================================================", flush=True)

    # 1. Health Check over HTTP
    req = urllib.request.Request(f"{BASE_URL}/health")
    with urllib.request.urlopen(req) as response:
        assert response.status == 200
        data = json.loads(response.read().decode())
        assert data["status"] == "ok"
        print("[PASS] HTTP Health Check Passed: ", data, flush=True)

    # 2. CORS Preflight Test
    req = urllib.request.Request(
        f"{BASE_URL}/api/resume/analyze",
        headers={
            "Origin": "http://127.0.0.1:5500",
            "Access-Control-Request-Method": "POST"
        },
        method="OPTIONS"
    )
    with urllib.request.urlopen(req) as response:
        assert response.status == 200
        origin_header = response.headers.get("access-control-allow-origin")
        assert origin_header in ["http://127.0.0.1:5500", "*"]
        print("[PASS] CORS Preflight Check Passed for origin http://127.0.0.1:5500", flush=True)

    # 3. Create DOCX Test File
    doc = Document()
    doc.add_heading("Akshay Live Test", level=0)
    doc.add_paragraph("Professional Title: Full Stack AI Developer")
    doc.add_paragraph("Summary: Experienced Python, FastAPI, and GenAI developer building scalable applications.")
    doc.add_paragraph("Skills: Python, FastAPI, SQL, Docker, LLMs, Git")
    doc.add_paragraph("Experience: Senior AI Engineer building LLM applications.")
    doc.add_paragraph("Education: Bachelor's Degree in Computer Science")

    test_docx_path = PROJECT_ROOT / "uploads" / "live_test_resume.docx"
    doc.save(str(test_docx_path))

    # 4. Test Live Multipart Upload & Analysis
    print("Sending live resume analysis request to backend...", flush=True)
    resume_res_json = post_multipart_file("/api/resume/analyze", test_docx_path, "live_test_resume.docx")
    assert resume_res_json["success"] is True
    resume_data = resume_res_json["analysis"]
    print("[PASS] Live AI Resume Analysis Endpoint Passed:", resume_data.get("candidate_name"), flush=True)

    # 5. Test Live JD Analysis
    jd_payload = {
        "job_description": "We are seeking a Senior AI Engineer skilled in Python, FastAPI, SQL, Docker, and Large Language Models."
    }
    jd_res_json = post_json("/api/jd/analyze", jd_payload)
    assert jd_res_json["success"] is True
    jd_data = jd_res_json["analysis"]
    print("[PASS] Live Job Description Analysis Endpoint Passed:", jd_data.get("job_title"), flush=True)

    # 6. Test Live Matching
    match_res_json = post_json("/api/match", {"resume_data": resume_data, "jd_data": jd_data})
    assert match_res_json["success"] is True
    match_data = match_res_json["match_result"]
    print("[PASS] Live Matching Engine Endpoint Passed. Score:", match_data.get("match_score"), flush=True)

    # 7. Test Live Tailoring
    tailor_res_json = post_json("/api/tailor", {"resume_data": resume_data, "jd_data": jd_data, "match_result": match_data})
    assert tailor_res_json["success"] is True
    tailor_data = tailor_res_json["tailoring"]
    print("[PASS] Live Resume Tailoring Endpoint Passed", flush=True)

    # 8. Test Live Generation
    gen_payload = {
        "resume_data": resume_data,
        "jd_data": jd_data,
        "match_result": match_data,
        "tailoring_result": tailor_data
    }
    gen_res_json = post_json("/api/generate", gen_payload)
    assert gen_res_json["success"] is True
    print("[PASS] Live Document Generation Endpoint Passed", flush=True)

    # 9. Test PDF Download
    pdf_content_type, pdf_bytes = get_binary(gen_res_json["download_urls"]["pdf"])
    assert pdf_content_type == "application/pdf"
    print("[PASS] Live PDF Download Endpoint Passed (Content-Length:", len(pdf_bytes), "bytes)", flush=True)

    # 10. Test DOCX Download
    docx_content_type, docx_bytes = get_binary(gen_res_json["download_urls"]["docx"])
    assert len(docx_bytes) > 0
    print("[PASS] Live DOCX Download Endpoint Passed (Content-Length:", len(docx_bytes), "bytes)", flush=True)

    print("==================================================", flush=True)
    print("ALL LIVE BACKEND ENDPOINTS & FRONTEND CONNECTIONS PASSED!", flush=True)
    print("==================================================", flush=True)

if __name__ == "__main__":
    test_live_backend()
