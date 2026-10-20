import os
import sys
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from fastapi.testclient import TestClient
from backend.main import app

client = TestClient(app)

def test_full_pipeline():
    print("==================================================", flush=True)
    print("Testing ResumeAI Full Pipeline (Phases 1-10)...", flush=True)
    print("==================================================", flush=True)

    # 1. Health Check
    response = client.get("/health")
    assert response.status_code == 200, f"Health check failed: {response.text}"
    print("[PASS] Health Check Passed", flush=True)

    # 2. Create valid test DOCX resume
    doc = Document()
    doc.add_heading("Akshay Developer", level=0)
    doc.add_paragraph("Professional Title: Senior Software Engineer")
    doc.add_paragraph("Summary: Experienced Python, FastAPI, and GenAI developer building scalable applications.")
    doc.add_paragraph("Skills: Python, FastAPI, SQL, Docker, LLMs, Git")
    doc.add_paragraph("Experience: Developed Python backend applications and integrated LLMs.")
    doc.add_paragraph("Education: Bachelor's Degree in Computer Science")

    sample_docx_path = PROJECT_ROOT / "uploads" / "sample_test_resume.docx"
    sample_docx_path.parent.mkdir(exist_ok=True)
    doc.save(str(sample_docx_path))

    # Test DOCX upload endpoint
    with open(sample_docx_path, "rb") as f:
        response = client.post(
            "/api/resume/upload",
            files={"resume": ("sample_test_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    assert response.status_code == 200, f"Resume upload failed: {response.text}"
    upload_result = response.json()
    assert upload_result["success"] is True
    print("[PASS] Resume Upload & Text Extraction Passed", flush=True)

    # 3. Analyze Resume with AI
    with open(sample_docx_path, "rb") as f:
        response = client.post(
            "/api/resume/analyze",
            files={"resume": ("sample_test_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    assert response.status_code == 200, f"Resume analysis failed: {response.text}"
    resume_analysis_result = response.json()
    assert resume_analysis_result["success"] is True
    resume_data = resume_analysis_result["analysis"]
    print("[PASS] AI Resume Analysis Passed for Candidate:", resume_data.get("candidate_name", "Akshay"), flush=True)

    # 4. Analyze Job Description
    sample_jd = """
    We are looking for a Senior AI Engineer with expertise in Python, FastAPI, SQL, Docker, and Large Language Models.
    Responsibilities include building REST APIs, prompt engineering, and integrating LLMs into modern web apps.
    Qualifications: Bachelor's degree in CS, 3+ years experience with Python and LLMs.
    """
    response = client.post(
        "/api/jd/analyze",
        json={"job_description": sample_jd}
    )
    assert response.status_code == 200, f"JD analysis failed: {response.text}"
    jd_analysis_result = response.json()
    assert jd_analysis_result["success"] is True
    jd_data = jd_analysis_result["analysis"]
    print("[PASS] Job Description Analysis Passed for Role:", jd_data.get("job_title", "Senior AI Engineer"), flush=True)

    # 5. Match Resume with JD
    response = client.post(
        "/api/match",
        json={"resume_data": resume_data, "jd_data": jd_data}
    )
    assert response.status_code == 200, f"Matching failed: {response.text}"
    match_result_payload = response.json()
    assert match_result_payload["success"] is True
    match_data = match_result_payload["match_result"]
    print("[PASS] Resume/JD Matching Passed. Compatibility Score:", match_data.get("match_score", match_data.get("skill_match_score")), flush=True)

    # 6. Resume Tailoring
    response = client.post(
        "/api/tailor",
        json={"resume_data": resume_data, "jd_data": jd_data, "match_result": match_data}
    )
    assert response.status_code == 200, f"Tailoring failed: {response.text}"
    tailor_result_payload = response.json()
    assert tailor_result_payload["success"] is True
    tailoring_data = tailor_result_payload["tailoring"]
    print("[PASS] Resume Tailoring Passed", flush=True)

    # 7. Generate Tailored Resume (LLM + PDF + DOCX)
    response = client.post(
        "/api/generate",
        json={
            "resume_data": resume_data,
            "jd_data": jd_data,
            "match_result": match_data,
            "tailoring_result": tailoring_data
        }
    )
    assert response.status_code == 200, f"Generation failed: {response.text}"
    gen_result_payload = response.json()
    assert gen_result_payload["success"] is True
    print("[PASS] Tailored Resume Generation Passed", flush=True)

    # 8. Download PDF
    pdf_download_url = gen_result_payload["download_urls"]["pdf"]
    response = client.get(pdf_download_url)
    assert response.status_code == 200, f"PDF Download failed: {response.status_code}"
    assert response.headers["content-type"] == "application/pdf"
    print("[PASS] PDF Download Endpoint Passed", flush=True)

    # 9. Download DOCX
    docx_download_url = gen_result_payload["download_urls"]["docx"]
    response = client.get(docx_download_url)
    assert response.status_code == 200, f"DOCX Download failed: {response.status_code}"
    print("[PASS] DOCX Download Endpoint Passed", flush=True)

    print("==================================================", flush=True)
    print("ALL PHASE 1-10 BACKEND & PIPELINE TESTS PASSED!", flush=True)
    print("==================================================", flush=True)

if __name__ == "__main__":
    test_full_pipeline()
