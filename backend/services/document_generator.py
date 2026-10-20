import html
from pathlib import Path
from docx import Document
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)


GENERATED_DIR = Path("generated_resumes")
GENERATED_DIR.mkdir(exist_ok=True)


def _safe_str(val) -> str:
    if isinstance(val, dict):
        parts = [_safe_str(v) for v in val.values() if v]
        return " | ".join(parts)
    if isinstance(val, list):
        return ", ".join(_safe_str(v) for v in val if v)
    return str(val)


def generate_docx(resume_data: dict, filename: str):

    file_path = GENERATED_DIR / filename

    document = Document()

    # Candidate Name
    candidate_name = resume_data.get("candidate_name", "")
    if candidate_name:
        heading = document.add_heading(_safe_str(candidate_name), level=0)
        heading.alignment = 1

    # Professional Title
    professional_title = resume_data.get("professional_title", "")
    if professional_title:
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        run = paragraph.add_run(_safe_str(professional_title))
        run.bold = True
        run.font.size = Pt(12)

    # Summary
    summary = resume_data.get("summary", "")
    if summary:
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(_safe_str(summary))

    # Skills
    skills = resume_data.get("skills", [])
    if skills:
        document.add_heading("Skills", level=1)
        skills_str = ", ".join([_safe_str(s) for s in skills]) if isinstance(skills, list) else _safe_str(skills)
        document.add_paragraph(skills_str)

    # Experience
    experience = resume_data.get("experience", [])
    if experience:
        document.add_heading("Experience", level=1)
        if isinstance(experience, list):
            for item in experience:
                document.add_paragraph(_safe_str(item), style="List Bullet")
        else:
            document.add_paragraph(_safe_str(experience))

    # Experience Summary
    experience_summary = resume_data.get("experience_summary", "")
    if experience_summary:
        document.add_heading("Experience Summary", level=1)
        document.add_paragraph(_safe_str(experience_summary))

    # Projects
    projects = resume_data.get("projects", [])
    if projects:
        document.add_heading("Projects", level=1)
        if isinstance(projects, list):
            for item in projects:
                document.add_paragraph(_safe_str(item), style="List Bullet")
        else:
            document.add_paragraph(_safe_str(projects))

    # Education
    education = resume_data.get("education", [])
    if education:
        document.add_heading("Education", level=1)
        if isinstance(education, list):
            for item in education:
                document.add_paragraph(_safe_str(item), style="List Bullet")
        else:
            document.add_paragraph(_safe_str(education))

    # Certifications
    certifications = resume_data.get("certifications", [])
    if certifications:
        document.add_heading("Certifications", level=1)
        if isinstance(certifications, list):
            for item in certifications:
                document.add_paragraph(_safe_str(item), style="List Bullet")
        else:
            document.add_paragraph(_safe_str(certifications))

    document.save(file_path)
    return str(file_path)


def generate_pdf(resume_data: dict, filename: str):

    file_path = GENERATED_DIR / filename

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    story = []

    # Candidate Name
    candidate_name = resume_data.get("candidate_name", "")
    if candidate_name:
        story.append(Paragraph(html.escape(_safe_str(candidate_name)), styles["Title"]))

    # Professional Title
    professional_title = resume_data.get("professional_title", "")
    if professional_title:
        story.append(Paragraph(html.escape(_safe_str(professional_title)), styles["Heading2"]))

    story.append(Spacer(1, 12))

    # Summary
    summary = resume_data.get("summary", "")
    if summary:
        story.append(Paragraph("Professional Summary", styles["Heading2"]))
        story.append(Paragraph(html.escape(_safe_str(summary)), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Skills
    skills = resume_data.get("skills", [])
    if skills:
        story.append(Paragraph("Skills", styles["Heading2"]))
        skills_str = ", ".join([_safe_str(s) for s in skills]) if isinstance(skills, list) else _safe_str(skills)
        story.append(Paragraph(html.escape(skills_str), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Experience
    experience = resume_data.get("experience", [])
    if experience:
        story.append(Paragraph("Experience", styles["Heading2"]))
        if isinstance(experience, list):
            for item in experience:
                story.append(Paragraph(f"&bull; {html.escape(_safe_str(item))}", styles["BodyText"]))
        else:
            story.append(Paragraph(html.escape(_safe_str(experience)), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Experience Summary
    experience_summary = resume_data.get("experience_summary", "")
    if experience_summary:
        story.append(Paragraph("Experience Summary", styles["Heading2"]))
        story.append(Paragraph(html.escape(_safe_str(experience_summary)), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Projects
    projects = resume_data.get("projects", [])
    if projects:
        story.append(Paragraph("Projects", styles["Heading2"]))
        if isinstance(projects, list):
            for item in projects:
                story.append(Paragraph(f"&bull; {html.escape(_safe_str(item))}", styles["BodyText"]))
        else:
            story.append(Paragraph(html.escape(_safe_str(projects)), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Education
    education = resume_data.get("education", [])
    if education:
        story.append(Paragraph("Education", styles["Heading2"]))
        if isinstance(education, list):
            for item in education:
                story.append(Paragraph(f"&bull; {html.escape(_safe_str(item))}", styles["BodyText"]))
        else:
            story.append(Paragraph(html.escape(_safe_str(education)), styles["BodyText"]))
        story.append(Spacer(1, 10))

    # Certifications
    certifications = resume_data.get("certifications", [])
    if certifications:
        story.append(Paragraph("Certifications", styles["Heading2"]))
        if isinstance(certifications, list):
            for item in certifications:
                story.append(Paragraph(f"&bull; {html.escape(_safe_str(item))}", styles["BodyText"]))
        else:
            story.append(Paragraph(html.escape(_safe_str(certifications)), styles["BodyText"]))

    document.build(story)
    return str(file_path)


if __name__ == "__main__":
    test_resume = {
        "candidate_name": "Akshay",
        "professional_title": "Python Developer",
        "summary": "Python developer with backend & GenAI experience <C++/FastAPI>.",
        "skills": ["Python", "FastAPI & SQL", "Git"],
        "experience": ["Developed Python backend applications."],
        "projects": ["AI Resume Analyzer & Matcher"],
        "education": ["Bachelor's Degree in Computer Science"],
        "certifications": []
    }

    pdf_path = generate_pdf(test_resume, "test_resume.pdf")
    print("PDF generated successfully:", pdf_path)
    docx_path = generate_docx(test_resume, "test_resume.docx")
    print("DOCX generated successfully:", docx_path)