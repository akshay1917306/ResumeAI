from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)

    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text and page_text.strip():
            text.append(page_text.strip())

    return "\n".join(text)


def extract_text_from_docx(file_path: str) -> str:
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        p_text = paragraph.text.strip()
        if p_text:
            text.append(p_text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    c_text = paragraph.text.strip()
                    if c_text and c_text not in text:
                        text.append(c_text)

    return "\n".join(text)


def extract_resume_text(file_path: str) -> str:
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        extracted = extract_text_from_pdf(file_path)

    elif extension == ".docx":
        extracted = extract_text_from_docx(file_path)

    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")

    if not extracted.strip():
        raise ValueError("No readable text was found in the uploaded file.")

    return extracted